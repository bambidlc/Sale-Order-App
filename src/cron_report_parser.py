"""Report parsing for scheduled Odoo imports, independent of the web converter API."""
import csv
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any, IO, Dict, List, Tuple, Union

import openpyxl
import xlrd


TEMPLATE_HEADERS = [
    "name",
    "partner_id/id",
    "user_id",
    "Cust #",
    "fecha",
    "order_line/product_uom_qty",
    "order_line/price_unit",
    "order_line/product_id",
]

_DEFAULT_SALESPERSON = "Jabes Omar De La Cruz"

_SALESPERSON_MAP: Dict[str, str] | None = None
_PARTNER_MAP: Dict[str, str] | None = None
_DESCRIPTIONS_MAP: Dict[str, str] | None = None
_REFERENCES_SET: set[str] | None = None

_DOC_KEYS = [
    "Header Document #",
    "Document #",
    "Doc #",
    "DOC #",
    "Doc#",
    "DOC#",
    "Doc No",
    "DOC NO",
    "Doc",
]
_CUSTOMER_KEYS = ["Customer Name", "CUSTOMER NAME", "Customer", "CLIENTE"]
_SKU_KEYS = ["Item Number", "SKU", "Sku", "sku", "Item", "ITEM", "Item #"]
_DESC_KEYS = [
    "Item Description",
    "Description",
    "DESCRIPTION",
    "description",
    "DESCRIPCION",
    "DESCRIPCIÃ“N",
]
_QTY_KEYS = ["Quantity", "Qty", "QTY"]
_PRICE_KEYS = ["Retail Price", "Price", "PRICE", "Unit Price", "UnitPrice"]
_STORE_KEYS = ["Store Number", "Store", "STORE"]
_PHONE_KEYS = ["Phone Number with Area Code", "Phone Number", "Phone", "PHONE"]
_CUST_KEYS = [
    "Customer Number",
    "Cust #",
    "Cust",
    "Customer #",
    "CUSTOMER #",
    "Customer ID",
    "CUSTOMER ID",
]
_DATE_KEYS = ["Date", "DATE", "date", "Fecha", "FECHA", "fecha", "Order Date", "ORDER DATE"]
_SALESPERSON_NAME_KEYS = ["Salesperson Name", "SALESPERSON NAME"]
_SALESPERSON_CODE_KEYS = [
    "Salesperson Number",
    "Salesperson",
    "SALESPERSON",
    "Sales Person",
    "Sales_Person",
    "Salesman",
    "Sales Man",
]
_TYPE_KEYS = ["Type", "TYPE", "type"]
_EMPTY_TOKENS = {"", "BLANK", "N/A", "NA", "NONE", "NULL"}


@dataclass(slots=True)
class OrderLineDraft:
    sku: str
    description: str
    quantity: float
    unit_price: float
    product_display: str


@dataclass(slots=True)
class SalesOrderDraft:
    order_name: str
    source_document: str
    store_number: str
    customer_number: str
    customer_name: str
    phone_number: str
    salesperson_name: str
    order_date_display: str
    order_date_odoo: str | None
    order_date_field_odoo: str | None
    partner_external_id: str | None
    lines: list[OrderLineDraft] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return bool(self.partner_external_id)

    @property
    def customer_label(self) -> str:
        return self.customer_name or self.customer_number or "Unknown"


def clear_caches() -> None:
    global _SALESPERSON_MAP, _PARTNER_MAP, _DESCRIPTIONS_MAP, _REFERENCES_SET
    _SALESPERSON_MAP = None
    _PARTNER_MAP = None
    _DESCRIPTIONS_MAP = None
    _REFERENCES_SET = None


def _normalize_header_name(value: str) -> str:
    return value.strip().casefold()


def _strip_quotes(value: str) -> str:
    stripped = value.strip()
    if len(stripped) >= 2 and stripped[0] == stripped[-1] and stripped[0] in {"'", '"'}:
        return stripped[1:-1]
    return stripped


def _load_descriptions_map(file_path: str = "Descriptions.csv") -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    if not os.path.exists(file_path):
        return mapping

    try:
        with open(file_path, "r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                default_code = (row.get("default_code") or "").strip()
                name = (row.get("name") or "").strip()
                if default_code and name:
                    mapping[default_code] = name
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                default_code = (row.get("default_code") or "").strip()
                name = (row.get("name") or "").strip()
                if default_code and name:
                    mapping[default_code] = name

    return mapping


def _load_salesperson_map(file_path: str = "Sales Person List.csv") -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    if not os.path.exists(file_path):
        return mapping

    try:
        with open(file_path, "r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                code = (row.get("Code") or "").strip().upper()
                name = (row.get("Name") or "").strip()
                if code and name:
                    mapping[code] = name
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                code = (row.get("Code") or "").strip().upper()
                name = (row.get("Name") or "").strip()
                if code and name:
                    mapping[code] = name

    return mapping


def _load_partner_map(file_path: str = "Contact (res.partner).csv") -> Dict[str, str]:
    mapping: Dict[str, str] = {}
    if not os.path.exists(file_path):
        return mapping

    def find_customer_number_key(fieldnames: List[str] | None) -> str | None:
        if not fieldnames:
            return None
        for fieldname in fieldnames:
            normalized = _normalize_header_name(fieldname)
            if "studio" in normalized and "cuenta" in normalized:
                return fieldname
        return None

    try:
        with open(file_path, "r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            account_key = find_customer_number_key(reader.fieldnames)
            if account_key:
                for row in reader:
                    customer_number = (row.get(account_key) or "").strip()
                    partner_id = (row.get("id") or "").strip()
                    if customer_number and partner_id:
                        mapping[customer_number] = partner_id
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            account_key = find_customer_number_key(reader.fieldnames)
            if account_key:
                for row in reader:
                    customer_number = (row.get(account_key) or "").strip()
                    partner_id = (row.get("id") or "").strip()
                    if customer_number and partner_id:
                        mapping[customer_number] = partner_id

    return mapping


def _load_references(file_path: str = "References.csv") -> set[str]:
    references: set[str] = set()
    if not os.path.exists(file_path):
        return references

    try:
        with open(file_path, "r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                reference = (row.get("Order Reference") or "").strip()
                if reference and len(reference) > 1:
                    references.add(reference[1:])
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=",")
            for row in reader:
                reference = (row.get("Order Reference") or "").strip()
                if reference and len(reference) > 1:
                    references.add(reference[1:])

    return references


def _get_salesperson_map() -> Dict[str, str]:
    global _SALESPERSON_MAP
    if _SALESPERSON_MAP is None:
        _SALESPERSON_MAP = _load_salesperson_map()
    return _SALESPERSON_MAP


def _get_partner_map() -> Dict[str, str]:
    global _PARTNER_MAP
    if _PARTNER_MAP is None:
        _PARTNER_MAP = _load_partner_map()
    return _PARTNER_MAP


def _get_descriptions_map() -> Dict[str, str]:
    global _DESCRIPTIONS_MAP
    if _DESCRIPTIONS_MAP is None:
        _DESCRIPTIONS_MAP = _load_descriptions_map()
    return _DESCRIPTIONS_MAP


def _get_references() -> set[str]:
    global _REFERENCES_SET
    if _REFERENCES_SET is None:
        _REFERENCES_SET = _load_references()
    return _REFERENCES_SET


def _coerce_decimal(raw_value: str) -> float:
    if not raw_value:
        return 0.0
    cleaned = raw_value.replace(",", "").strip()
    try:
        return float(cleaned)
    except Exception:
        return 0.0


def _parse_date_value(date_str: str) -> datetime | None:
    if not date_str:
        return None

    try:
        serial_date = float(date_str)
        if 1 <= serial_date <= 100000:
            excel_epoch = datetime(1899, 12, 30)
            return excel_epoch + timedelta(days=serial_date)
    except (ValueError, TypeError):
        pass

    formats = [
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%m/%d/%Y %H:%M:%S.%f",
        "%d/%m/%Y %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%m-%d-%Y",
        "%d/%m/%y",
        "%m/%d/%y",
    ]
    date_str_clean = str(date_str).strip()
    for fmt in formats:
        try:
            return datetime.strptime(date_str_clean, fmt)
        except ValueError:
            continue

    if " " in date_str_clean:
        date_only = date_str_clean.split(" ")[0]
        for fmt in formats:
            try:
                return datetime.strptime(date_only, fmt)
            except ValueError:
                continue

    if "T" in date_str_clean:
        date_only = date_str_clean.split("T")[0]
        for fmt in formats:
            try:
                return datetime.strptime(date_only, fmt)
            except ValueError:
                continue

    return None


def _format_template_date(date_str: str) -> str:
    parsed = _parse_date_value(date_str)
    if parsed is None:
        return date_str.strip() if date_str else ""
    return parsed.strftime("%m/%d/%Y")


def _format_odoo_date(date_str: str) -> str | None:
    parsed = _parse_date_value(date_str)
    if parsed is None:
        return None
    return parsed.strftime("%Y-%m-%d %H:%M:%S")


def _format_odoo_date_field(date_str: str) -> str | None:
    parsed = _parse_date_value(date_str)
    if parsed is None:
        return None
    return parsed.strftime("%Y-%m-%d")


def _resolve_headers(row: Dict[str, Any]) -> Dict[str, str]:
    return {_normalize_header_name(key): key for key in row.keys()}


def _get_value(row: Dict[str, Any], keys: List[str]) -> str:
    for key in keys:
        if key in row and row[key] is not None:
            return str(row.get(key, "")).strip()

    headers = _resolve_headers(row)
    for key in keys:
        actual = headers.get(_normalize_header_name(key))
        if actual and row.get(actual) is not None:
            return str(row.get(actual, "")).strip()

    return ""


def _find_header_key(header_keys: List[str], keys: List[str]) -> str | None:
    normalized = {_normalize_header_name(key): key for key in header_keys}
    for key in keys:
        actual = normalized.get(_normalize_header_name(key))
        if actual:
            return actual
    return None


def _build_order_name(document_number: str, type_value: str) -> str:
    clean_document = re.sub(r"\s+", "", document_number)
    if clean_document:
        return clean_document
    return f"Q{datetime.now().strftime('%m%d%H%M')}"


def _resolve_salesperson_name(row: Dict[str, Any], salesperson_map: Dict[str, str]) -> str:
    salesperson_name = _get_value(row, _SALESPERSON_NAME_KEYS)
    if salesperson_name and salesperson_name.strip().upper() not in _EMPTY_TOKENS:
        return salesperson_name

    salesperson_code = _get_value(row, _SALESPERSON_CODE_KEYS).upper()
    if salesperson_code and salesperson_code not in _EMPTY_TOKENS:
        return salesperson_map.get(salesperson_code, salesperson_code)

    return ""


def _resolve_product_display(sku: str, description: str, descriptions_map: Dict[str, str]) -> str:
    mapped_name = descriptions_map.get(sku)
    if mapped_name:
        return f"[{sku}] {mapped_name}"
    return f"[{sku}] {description}"


def _hydrate_order_metadata(
    draft: SalesOrderDraft,
    row: Dict[str, Any],
    customer_key: str | None,
    date_key: str | None,
    salesperson_map: Dict[str, str],
    partner_map: Dict[str, str],
) -> None:
    if not draft.customer_number:
        draft.customer_number = _get_value(row, _CUST_KEYS)
        if draft.customer_number and not draft.partner_external_id:
            draft.partner_external_id = partner_map.get(draft.customer_number)

    if not draft.store_number:
        draft.store_number = _get_value(row, _STORE_KEYS)

    if not draft.customer_name:
        draft.customer_name = _get_value(row, [customer_key]) if customer_key else _get_value(row, _CUSTOMER_KEYS)

    if not draft.phone_number:
        draft.phone_number = _get_value(row, _PHONE_KEYS)

    if not draft.salesperson_name or draft.salesperson_name == _DEFAULT_SALESPERSON:
        draft.salesperson_name = _resolve_salesperson_name(row, salesperson_map)

    if not draft.order_date_display:
        raw_date = _get_value(row, [date_key]) if date_key else _get_value(row, _DATE_KEYS)
        draft.order_date_display = _format_template_date(raw_date)
        draft.order_date_odoo = _format_odoo_date(raw_date)
        draft.order_date_field_odoo = _format_odoo_date_field(raw_date)


def build_sales_order_drafts(
    epicor_rows: List[Dict[str, Any]],
    *,
    skip_existing_references: bool = False,
    resolve_local_partner_ids: bool = True,
    resolve_local_product_names: bool = True,
) -> List[SalesOrderDraft]:
    if not epicor_rows:
        return []

    header_keys = list(epicor_rows[0].keys())
    doc_key = _find_header_key(header_keys, _DOC_KEYS)
    customer_key = _find_header_key(header_keys, _CUSTOMER_KEYS)
    date_key = _find_header_key(header_keys, _DATE_KEYS)

    partner_map = _get_partner_map() if resolve_local_partner_ids else {}
    descriptions_map = _get_descriptions_map() if resolve_local_product_names else {}
    salesperson_map = _get_salesperson_map()
    references = _get_references() if skip_existing_references else set()

    drafts_by_key: Dict[str, SalesOrderDraft] = {}

    for row in epicor_rows:
        sku = _get_value(row, _SKU_KEYS)
        if not sku:
            continue

        description = _get_value(row, _DESC_KEYS)
        sku_lower = sku.lower()
        desc_lower = description.lower()
        if "total" in sku_lower or "total" in desc_lower:
            continue

        document_number = _get_value(row, [doc_key]) if doc_key else ""
        if doc_key and (not document_number or document_number == "0"):
            continue

        if skip_existing_references and document_number and document_number in references:
            continue

        order_key = document_number or "__single_order__"
        draft = drafts_by_key.get(order_key)
        if draft is None:
            customer_number = _get_value(row, _CUST_KEYS)
            store_number = _get_value(row, _STORE_KEYS)
            customer_name = _get_value(row, [customer_key]) if customer_key else _get_value(row, _CUSTOMER_KEYS)
            phone_number = _get_value(row, _PHONE_KEYS)
            salesperson_name = _resolve_salesperson_name(row, salesperson_map)
            raw_date = _get_value(row, [date_key]) if date_key else _get_value(row, _DATE_KEYS)
            partner_external_id = partner_map.get(customer_number) if customer_number else None
            type_value = _get_value(row, _TYPE_KEYS)
            order_name = _build_order_name(document_number, type_value)
            draft = SalesOrderDraft(
                order_name=order_name,
                source_document=document_number,
                store_number=store_number,
                customer_number=customer_number,
                customer_name=customer_name,
                phone_number=phone_number,
                salesperson_name=salesperson_name,
                order_date_display=_format_template_date(raw_date),
                order_date_odoo=_format_odoo_date(raw_date),
                order_date_field_odoo=_format_odoo_date_field(raw_date),
                partner_external_id=partner_external_id,
            )
            drafts_by_key[order_key] = draft
        else:
            _hydrate_order_metadata(draft, row, customer_key, date_key, salesperson_map, partner_map)

        quantity = _coerce_decimal(_get_value(row, _QTY_KEYS))
        unit_price = _coerce_decimal(_get_value(row, _PRICE_KEYS))
        draft.lines.append(
            OrderLineDraft(
                sku=sku,
                description=description,
                quantity=quantity,
                unit_price=unit_price,
                product_display=_resolve_product_display(sku, description, descriptions_map),
            )
        )

    return [draft for draft in drafts_by_key.values() if draft.lines]


def parse_rows(file_path: str) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    _, ext = os.path.splitext(file_path)

    if ext.lower() == ".csv":
        try:
            with open(file_path, "r", encoding="utf-8", newline="") as handle:
                first_line = handle.readline()
                delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1", newline="") as handle:
                first_line = handle.readline()
                delimiter = ";" if first_line.count(";") > first_line.count(",") else ","

        try:
            with open(file_path, "r", encoding="utf-8", newline="") as handle:
                for row in csv.DictReader(handle, delimiter=delimiter):
                    rows.append(row)
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1", newline="") as handle:
                for row in csv.DictReader(handle, delimiter=delimiter):
                    rows.append(row)
    elif ext.lower() == ".xlsx":
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        for row_index in range(2, sheet.max_row + 1):
            row_dict: Dict[str, Any] = {}
            for column_index, header in enumerate(headers):
                cell_value = sheet.cell(row=row_index, column=column_index + 1).value
                row_dict[str(header)] = str(cell_value) if cell_value is not None else ""
            rows.append(row_dict)
    elif ext.lower() == ".xls":
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)
        headers = [sheet.cell_value(0, column) for column in range(sheet.ncols)]
        for row_index in range(1, sheet.nrows):
            row_dict: Dict[str, Any] = {}
            for column_index, header in enumerate(headers):
                cell_value = sheet.cell_value(row_index, column_index)
                row_dict[str(header)] = str(cell_value) if cell_value is not None else ""
            rows.append(row_dict)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    return rows
